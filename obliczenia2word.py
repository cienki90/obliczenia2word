#!/usr/bin/env python3
"""
obliczenia2word – automatyczne wstawianie tabel z obliczeniami (Excel) do projektów (Word).

Uruchomiony bez parametrów (np. dwuklikiem w plik .exe) program pyta w okienkach o:
  1. plik Excel z obliczeniami,
  2. folder z opisami (.docx),
  3. folder, w którym zapisać gotowe opisy.

Dla każdego opisu (np. 181.docx) program:
  * wybiera arkusz o tej samej nazwie/numerze (np. arkusz "181"),
  * odtwarza zakres wydruku arkusza jako tabelę Word z formatowaniem z Excela,
  * wstawia ją w rozdziale „Zestawienie zbiorcze słupów” (jak w 112.docx i 115.docx),
  * pilnuje, żeby za tabelą nie powstała pusta strona.
Na koniec pokazuje raport, w tym listę czerwonych komórek w tabelach z obliczeniami.

Wiersz poleceń (opcjonalnie):
    python obliczenia2word.py -x "Obliczenia.xlsx" -o wynik folder_z_opisami
    python obliczenia2word.py --arkusz 181 projekt.docx
    python obliczenia2word.py --nadpisz 181.docx
"""
from __future__ import annotations

import argparse
import glob
import math
import os
import re
import sys
from xml.sax.saxutils import escape

try:
    import openpyxl
    from openpyxl.utils import range_boundaries
    import docx
    from docx.oxml import parse_xml
except ImportError:  # pragma: no cover
    print("Brak wymaganych bibliotek. Zainstaluj je poleceniem:\n"
          "    pip install python-docx openpyxl")
    sys.exit(1)

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = "{%s}" % W_NS

APP_NAME = "Obliczenia do Worda"
HEADING_TEXT = "Zestawienie zbiorcze słupów"
DEFAULT_EMPTY_PARAGRAPHS_BEFORE = 5   # tak jak w 112.docx i 115.docx
KEEP_WITH_NEXT = re.compile(r"obw nr|\bLp\b", re.IGNORECASE)

# Excel styl obramowania -> (styl Word, grubość w 1/8 pt)
BORDER_MAP = {
    "hair": ("single", 2),
    "thin": ("single", 4),
    "medium": ("single", 12),
    "thick": ("single", 18),
    "double": ("double", 4),
    "dotted": ("dotted", 4),
    "dashed": ("dashed", 4),
    "mediumDashed": ("dashed", 12),
    "dashDot": ("dotDash", 4),
    "mediumDashDot": ("dotDash", 12),
    "dashDotDot": ("dotDotDash", 4),
    "mediumDashDotDot": ("dotDotDash", 12),
    "slantDashDot": ("dashDotStroked", 12),
}

H_ALIGN_MAP = {
    "left": "left",
    "center": "center",
    "centerContinuous": "center",
    "right": "right",
    "justify": "both",
    "distributed": "distribute",
    "fill": "left",
}

V_ALIGN_MAP = {"top": "top", "center": "center", "bottom": "bottom",
               "justify": "center", "distributed": "center"}


# --------------------------------------------------------------------------- #
#  Excel -> model tabeli
# --------------------------------------------------------------------------- #
THEME_COLORS: list[str] = []   # kolory motywu bieżącego skoroszytu (indeks jak w Excelu)


def load_theme_colors(wb) -> list[str]:
    """Odczytuje paletę motywu skoroszytu (potrzebną dla kolorów typu „motyw”)."""
    order = ["lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3",
             "accent4", "accent5", "accent6", "hlink", "folHlink"]
    colors = {}
    try:
        from lxml import etree
        root = etree.fromstring(wb.loaded_theme)
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        scheme = root.find(".//a:clrScheme", ns)
        for el in scheme:
            name = etree.QName(el).localname
            clr = el[0]
            colors[name] = (clr.get("val") if etree.QName(clr).localname == "srgbClr"
                            else clr.get("lastClr")) or None
    except Exception:
        pass
    defaults = ["FFFFFF", "000000", "E7E6E6", "44546A", "4472C4", "ED7D31",
                "A5A5A5", "FFC000", "5B9BD5", "70AD47", "0563C1", "954F72"]
    return [(colors.get(n) or d).upper() for n, d in zip(order, defaults)]


def _apply_tint(rgb: str, tint: float) -> str:
    import colorsys
    r, g, b = (int(rgb[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, lum, s = colorsys.rgb_to_hls(r, g, b)
    lum = lum * (1 + tint) if tint < 0 else lum * (1 - tint) + tint
    r, g, b = colorsys.hls_to_rgb(h, max(0.0, min(1.0, lum)), s)
    return "%02X%02X%02X" % (round(r * 255), round(g * 255), round(b * 255))


def _rgb(color) -> str | None:
    """Zwraca kolor w formacie RRGGBB (obsługuje kolory RGB, z motywu i indeksowane)."""
    if color is None:
        return None
    try:
        ctype = color.type
        if ctype == "rgb":
            rgb = color.rgb
            if not isinstance(rgb, str) or len(rgb) < 6:
                return None
            rgb = rgb[-6:].upper()
        elif ctype == "theme":
            idx = int(color.theme)
            palette = THEME_COLORS or load_theme_colors(None)
            if not 0 <= idx < len(palette):
                return None
            rgb = palette[idx]
        elif ctype == "indexed":
            from openpyxl.styles.colors import COLOR_INDEX
            idx = int(color.indexed)
            if idx >= 64 or idx >= len(COLOR_INDEX):   # kolory systemowe
                return None
            rgb = COLOR_INDEX[idx][-6:].upper()
        else:
            return None
        tint = float(color.tint or 0)
        return _apply_tint(rgb, tint) if tint else rgb
    except Exception:
        return None


def is_red(rgb: str | None) -> bool:
    """Czy kolor jest czerwony (od jasnoczerwonego/różowego tła po ciemną czerwień)."""
    if not rgb:
        return False
    import colorsys
    r, g, b = (int(rgb[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, s, v = colorsys.rgb_to_hsv(r, g, b)
    hue = h * 360
    return (hue <= 15 or hue >= 340) and s >= 0.15 and v >= 0.35


def find_red_cells(ws, ws_values):
    """Zwraca listę czerwonych komórek w tabeli (zakresie wydruku) arkusza:
    [(adres, opis, wartość), ...]."""
    min_col, min_row, max_col, max_row = sheet_area(ws)
    covered = set()
    for rng in ws.merged_cells.ranges:
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if (r, c) != (rng.min_row, rng.min_col):
                    covered.add((r, c))
    found = []
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            if (r, c) in covered:
                continue
            cell = ws.cell(r, c)
            value = ws_values.cell(r, c).value
            reasons = []
            if cell.fill is not None and cell.fill.fill_type and is_red(_rgb(cell.fill.fgColor)):
                reasons.append("czerwone tło")
            if value not in (None, "") and is_red(_rgb(cell.font.color)):
                reasons.append("czerwony tekst")
            if reasons:
                text = format_value(value, cell.number_format).replace("\n", " ")
                found.append((cell.coordinate, ", ".join(reasons), text))
    return found


def format_value(value, number_format: str) -> str:
    """Formatuje wartość komórki tak, jak wyświetla ją Excel (dla typowych formatów)."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "PRAWDA" if value else "FAŁSZ"
    if isinstance(value, (int, float)):
        fmt = (number_format or "General").split(";")[0]
        if fmt == "@":
            return str(value)
        m = re.fullmatch(r'[#,]*0(?:\.(0+))?(%?)', fmt.replace("\\", "").replace('"', ""))
        if m:
            decimals = len(m.group(1) or "")
            v = value * 100 if m.group(2) else value
            thousands = "," in fmt
            s = f"{v:,.{decimals}f}" if thousands else f"{v:.{decimals}f}"
            if thousands:
                s = s.replace(",", "\u00a0")
            return s + m.group(2)
        # General
        if isinstance(value, float):
            if math.isfinite(value) and value == int(value) and abs(value) < 1e15:
                return str(int(value))
            return f"{value:.10g}"
        return str(value)
    if hasattr(value, "strftime"):
        return value.strftime("%d.%m.%Y")
    return str(value)


def column_widths_chars(ws, min_col: int, max_col: int) -> list[float]:
    default = ws.sheet_format.defaultColWidth or ((ws.sheet_format.baseColWidth or 8) + 0.71)
    widths = {c: default for c in range(min_col, max_col + 1)}
    hidden = set()
    for dim in ws.column_dimensions.values():
        lo, hi = dim.min or 0, dim.max or 0
        for c in range(lo, hi + 1):
            if c in widths:
                if dim.width:
                    widths[c] = dim.width
                if dim.hidden:
                    hidden.add(c)
    return [0.0 if c in hidden else widths[c] for c in range(min_col, max_col + 1)]


def sheet_area(ws):
    """Zakres tabeli: obszar wydruku arkusza, a gdy go nie ma – używany zakres."""
    area = None
    if ws.print_area:
        try:
            pa = ws.print_area
            pa = pa[0] if isinstance(pa, (list, tuple)) else pa
            pa = pa.split(",")[0].split("!")[-1].replace("$", "")
            area = range_boundaries(pa)
        except Exception:
            area = None
    if not area or None in area:
        area = range_boundaries(ws.calculate_dimension())
    return area


def read_sheet(ws, ws_values):
    """Czyta arkusz i zwraca prosty model: kolumny, wiersze, komórki z formatowaniem."""
    min_col, min_row, max_col, max_row = sheet_area(ws)

    col_chars = column_widths_chars(ws, min_col, max_col)
    default_h = ws.sheet_format.defaultRowHeight or 15

    merged_at = {}
    covered = set()
    for rng in ws.merged_cells.ranges:
        if rng.max_row < min_row or rng.min_row > max_row or rng.max_col < min_col or rng.min_col > max_col:
            continue
        merged_at[(rng.min_row, rng.min_col)] = rng
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if (r, c) != (rng.min_row, rng.min_col):
                    covered.add((r, c))

    visible_cols = [c for i, c in enumerate(range(min_col, max_col + 1)) if col_chars[i] > 0]
    rows = []
    for r in range(min_row, max_row + 1):
        dim = ws.row_dimensions.get(r)
        if dim is not None and dim.hidden:
            continue
        height = (dim.height if dim is not None and dim.height else None) or default_h
        cells = []
        for c in visible_cols:
            if (r, c) in covered:
                rng = next((m for (mr, mc), m in merged_at.items()
                            if m.min_row < r <= m.max_row and m.min_col <= c <= m.max_col
                            and m.min_col == c), None)
                if rng is not None:  # kontynuacja scalenia pionowego
                    cells.append({"vmerge_continue": True, "col": c,
                                  "span": _visible_span(visible_cols, rng.min_col, rng.max_col),
                                  "borders": _merged_borders(ws, rng)})
                continue
            rng = merged_at.get((r, c))
            cell = ws.cell(r, c)
            value = ws_values.cell(r, c).value
            if isinstance(value, str) and value.startswith("=") and isinstance(cell.value, str):
                value = None  # formuła bez zapisanego wyniku
            if rng is not None:
                span = _visible_span(visible_cols, rng.min_col, rng.max_col)
                borders = _merged_borders(ws, rng)
                vmerge = rng.max_row > rng.min_row
            else:
                span, borders, vmerge = 1, _cell_borders(cell), False
            cells.append({
                "col": c,
                "span": span,
                "vmerge_start": vmerge,
                "borders": borders,
                "text": format_value(value, cell.number_format),
                "is_number": isinstance(value, (int, float)) and not isinstance(value, bool),
                "font": cell.font,
                "alignment": cell.alignment,
                "fill": cell.fill,
            })
        rows.append({"height": height, "cells": cells})
    widths = [col_chars[c - min_col] for c in visible_cols]
    return widths, rows, ws_values.title


def _visible_span(visible_cols, c1, c2):
    return max(1, sum(1 for c in visible_cols if c1 <= c <= c2))


def _side(border_side):
    if border_side is None or not border_side.style:
        return None
    style, size = BORDER_MAP.get(border_side.style, ("single", 4))
    return style, size, _rgb(border_side.color) or "auto"


def _cell_borders(cell):
    b = cell.border
    return {"top": _side(b.top), "left": _side(b.left),
            "bottom": _side(b.bottom), "right": _side(b.right)}


def _merged_borders(ws, rng):
    def first(sides):
        return next((s for s in sides if s), None)
    rows = range(rng.min_row, rng.max_row + 1)
    cols = range(rng.min_col, rng.max_col + 1)
    return {
        "top": first(_side(ws.cell(rng.min_row, c).border.top) for c in cols),
        "bottom": first(_side(ws.cell(rng.max_row, c).border.bottom) for c in cols),
        "left": first(_side(ws.cell(r, rng.min_col).border.left) for r in rows),
        "right": first(_side(ws.cell(r, rng.max_col).border.right) for r in rows),
    }


# --------------------------------------------------------------------------- #
#  Szerokości kolumn
# --------------------------------------------------------------------------- #
def _text_twips(text, font):
    """Przybliżona szerokość tekstu (twips) dla czcionki z Excela."""
    size_pt = float(font.sz or 11)
    factor = 0.60 if font.b else 0.55
    return len(text) * size_pt * factor * 20


def compute_grid(widths_chars, rows, available_twips):
    """Szerokości kolumn liczone jak autodopasowanie Worda: każda kolumna dostaje co
    najmniej szerokość najdłuższego słowa, a pozostałe miejsce jest dzielone
    proporcjonalnie do tego, ile kolumna potrzebuje, by tekst się nie zawijał.
    Kolumny bez tekstu dostają szerokość proporcjonalną do szerokości w Excelu."""
    n = len(widths_chars)
    cell_margin = 140 + 30  # marginesy komórki (2 x 70) + zapas
    minimum = [0.0] * n
    maximum = [0.0] * n
    for row in rows:
        idx = 0
        for cell in row["cells"]:
            span = cell["span"]
            text = cell.get("text")
            if span == 1 and text and idx < n:
                font = cell["font"]
                words = re.split(r"\s+", text)
                lines = text.split("\n")
                minimum[idx] = max(minimum[idx], max(_text_twips(w, font) for w in words) + cell_margin)
                maximum[idx] = max(maximum[idx], max(_text_twips(ln, font) for ln in lines) + cell_margin)
            idx += span

    px = [int(w * 7 + 5) if w > 0 else 0 for w in widths_chars]
    excel = [p * available_twips / (sum(px) or 1) for p in px]
    # kolumny bez własnego tekstu: minimum wg Excela (ale nie więcej niż potrzeba)
    for i in range(n):
        if minimum[i] == 0 and px[i] > 0:
            minimum[i] = maximum[i] = min(excel[i], 400)

    total_min = sum(minimum)
    if total_min >= available_twips:  # nie mieści się – skalujemy minimum
        grid = [m * available_twips / total_min for m in minimum]
    else:
        extra = available_twips - total_min
        want = [max(0.0, mx - mn) for mx, mn in zip(maximum, minimum)]
        total_want = sum(want)
        if total_want <= extra:  # wszystko mieści się w jednej linii
            spare = extra - total_want
            grid = [mx + spare * e / (sum(excel) or 1) for mx, e in zip(maximum, excel)]
        else:
            grid = [mn + extra * w / total_want for mn, w in zip(minimum, want)]
    grid = [int(round(g)) for g in grid]
    grid[-1] += available_twips - sum(grid)
    return grid


# --------------------------------------------------------------------------- #
#  Model -> XML tabeli Word
# --------------------------------------------------------------------------- #
def _border_xml(name, side):
    if side is None:
        return f'<w:{name} w:val="nil"/>'
    style, size, color = side
    return f'<w:{name} w:val="{style}" w:sz="{size}" w:space="0" w:color="{color}"/>'


def _rpr_xml(font):
    name = escape(font.name or "Calibri", {'"': "&quot;"})
    parts = [f'<w:rFonts w:ascii="{name}" w:eastAsia="Times New Roman" w:hAnsi="{name}" w:cs="{name}"/>']
    if font.b:
        parts.append("<w:b/><w:bCs/>")
    if font.i:
        parts.append("<w:i/><w:iCs/>")
    if font.strike:
        parts.append("<w:strike/>")
    color = _rgb(font.color)
    if color and color != "000000":
        parts.append(f'<w:color w:val="{color}"/>')
    if font.sz:
        hp = int(round(float(font.sz) * 2))
        parts.append(f'<w:sz w:val="{hp}"/><w:szCs w:val="{hp}"/>')
    if font.u:
        parts.append('<w:u w:val="{}"/>'.format("double" if "double" in str(font.u) else "single"))
    if font.vertAlign in ("superscript", "subscript"):
        parts.append(f'<w:vertAlign w:val="{font.vertAlign}"/>')
    parts.append('<w:lang w:val="pl-PL"/>')
    return "<w:rPr>" + "".join(parts) + "</w:rPr>"


def _runs_xml(text, rpr):
    out = []
    for i, line in enumerate(text.replace("\r\n", "\n").split("\n")):
        if i:
            out.append(f"<w:r>{rpr}<w:br/></w:r>")
        if line:
            out.append(f'<w:r>{rpr}<w:t xml:space="preserve">{escape(line)}</w:t></w:r>')
    return "".join(out)


def build_table_xml(grid, rows):
    xml = [f'<w:tbl xmlns:w="{W_NS}">',
           "<w:tblPr>",
           f'<w:tblW w:w="{sum(grid)}" w:type="dxa"/>',
           '<w:jc w:val="center"/>',
           '<w:tblLayout w:type="autofit"/>',
           '<w:tblCellMar><w:left w:w="70" w:type="dxa"/><w:right w:w="70" w:type="dxa"/></w:tblCellMar>',
           '<w:tblLook w:val="04A0" w:firstRow="1" w:lastRow="0" w:firstColumn="1" '
           'w:lastColumn="0" w:noHBand="0" w:noVBand="1"/>',
           "</w:tblPr>",
           "<w:tblGrid>" + "".join(f'<w:gridCol w:w="{g}"/>' for g in grid) + "</w:tblGrid>"]

    for row in rows:
        height = int(round(row["height"] * 20))
        row_text = " ".join(c.get("text", "") for c in row["cells"])
        # nagłówek obwodu i nagłówek kolumn trzymamy razem z kolejnym wierszem,
        # żeby nie zostały same na dole strony
        keep_next = "<w:keepNext/>" if KEEP_WITH_NEXT.search(row_text) else ""
        xml.append(f'<w:tr><w:trPr><w:cantSplit/><w:trHeight w:val="{height}"/></w:trPr>')
        idx = 0
        for cell in row["cells"]:
            span = cell["span"]
            width = sum(grid[idx:idx + span])
            idx += span
            b = cell["borders"]
            tcpr = [f'<w:tcW w:w="{width}" w:type="dxa"/>']
            if span > 1:
                tcpr.append(f'<w:gridSpan w:val="{span}"/>')
            if cell.get("vmerge_continue"):
                tcpr.append("<w:vMerge/>")
            elif cell.get("vmerge_start"):
                tcpr.append('<w:vMerge w:val="restart"/>')
            tcpr.append("<w:tcBorders>" + "".join(_border_xml(n, b[n]) for n in
                                                  ("top", "left", "bottom", "right")) + "</w:tcBorders>")
            if cell.get("vmerge_continue"):
                xml.append("<w:tc><w:tcPr>" + "".join(tcpr) + "</w:tcPr><w:p/></w:tc>")
                continue

            fill = cell["fill"]
            if fill is not None and fill.fill_type == "solid":
                color = _rgb(fill.fgColor)
                if color:
                    tcpr.append(f'<w:shd w:val="clear" w:color="auto" w:fill="{color}"/>')
            al = cell["alignment"]
            tcpr.append(f'<w:vAlign w:val="{V_ALIGN_MAP.get(al.vertical or "bottom", "bottom")}"/>')

            if al.horizontal in H_ALIGN_MAP:
                jc = H_ALIGN_MAP[al.horizontal]
            else:  # "general": liczby do prawej, tekst do lewej
                jc = "right" if cell["is_number"] else "left"
            rpr = _rpr_xml(cell["font"])
            ppr = f'<w:pPr>{keep_next}<w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>' \
                  f'<w:jc w:val="{jc}"/>{rpr}</w:pPr>'
            xml.append("<w:tc><w:tcPr>" + "".join(tcpr) + "</w:tcPr><w:p>" + ppr
                       + _runs_xml(cell["text"], rpr) + "</w:p></w:tc>")
        xml.append("</w:tr>")
    xml.append("</w:tbl>")
    return "".join(xml)


# --------------------------------------------------------------------------- #
#  Wstawianie do dokumentu
# --------------------------------------------------------------------------- #
def _text(el):
    return "".join(t.text or "" for t in el.iter(W + "t"))


def _is_empty_paragraph(el):
    return (el.tag == W + "p" and not _text(el).strip()
            and el.find(".//" + W + "drawing") is None
            and el.find(".//" + W + "pict") is None
            and el.find(".//" + W + "sectPr") is None)


def find_insertion_point(document, empty_before):
    body = document.element.body
    children = list(body.iterchildren())
    heading_idx = None
    for i, el in enumerate(children):
        if el.tag == W + "p" and _text(el).strip().rstrip(":").lower() == HEADING_TEXT.lower():
            heading_idx = i
    if heading_idx is None:
        raise ValueError(f'nie znaleziono nagłówka "{HEADING_TEXT}"')

    # istniejąca tabela (dokument już wypełniony) -> do podmiany
    anchor = children[heading_idx]
    existing = None
    empties = 0
    for el in children[heading_idx + 1:]:
        if el.tag == W + "tbl":
            existing = el
            break
        if not _is_empty_paragraph(el):
            break
        empties += 1
        if empties <= empty_before:
            anchor = el
    return anchor, existing


def section_text_width(document, anchor):
    """Szerokość obszaru tekstu (twips) sekcji, w której znajdzie się tabela."""
    sect = None
    el = anchor.getnext()
    while el is not None:
        sect = el.find(".//" + W + "sectPr") if el.tag == W + "p" else (el if el.tag == W + "sectPr" else None)
        if sect is not None:
            break
        el = el.getnext()
    if sect is None:
        sect = document.element.body.find(W + "sectPr")
    pg, mar = sect.find(W + "pgSz"), sect.find(W + "pgMar")
    width = int(pg.get(W + "w")) - int(mar.get(W + "left")) - int(mar.get(W + "right"))
    if mar.get(W + "gutter"):
        width -= int(mar.get(W + "gutter"))
    return width


def sheet_for_document(wb, docx_path, sheet_name=None):
    if sheet_name:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f'w pliku Excel nie ma arkusza "{sheet_name}"')
        return sheet_name
    stem = os.path.splitext(os.path.basename(docx_path))[0]
    if stem in wb.sheetnames:
        return stem
    numbers = re.findall(r"\d+", stem)
    for num in numbers:
        for name in wb.sheetnames:
            if re.fullmatch(r"\d+", name.strip()) and int(name) == int(num):
                return name
    raise ValueError(f'nie znaleziono arkusza pasującego do pliku "{os.path.basename(docx_path)}"')


# --------------------------------------------------------------------------- #
#  Pilnowanie, żeby za tabelą nie powstała pusta strona
# --------------------------------------------------------------------------- #
def _has_page_break(p):
    return any(br.get(W + "type") == "page" for br in p.iter(W + "br"))


def _has_page_break_before(p):
    ppr = p.find(W + "pPr")
    pb = ppr.find(W + "pageBreakBefore") if ppr is not None else None
    return pb is not None and pb.get(W + "val") not in ("0", "false", "off")


def _set_page_break_before(p):
    ppr = p.find(W + "pPr")
    if ppr is None:
        ppr = parse_xml(f'<w:pPr xmlns:w="{W_NS}"/>')
        p.insert(0, ppr)
    old = ppr.find(W + "pageBreakBefore")
    if old is not None:
        ppr.remove(old)
    pb = parse_xml(f'<w:pageBreakBefore xmlns:w="{W_NS}"/>')
    # kolejność wg schematu: pStyle, keepNext, keepLines, pageBreakBefore, ...
    pos = 0
    for i, child in enumerate(ppr):
        if child.tag in (W + "pStyle", W + "keepNext", W + "keepLines"):
            pos = i + 1
    ppr.insert(pos, pb)


def prevent_blank_page(table):
    """Za tabelą w szablonie są puste akapity i akapit ze znakiem podziału strony.
    Gdy tabela kończy się na dole strony, puste akapity przelewają się na następną
    stronę, a podział strony tworzy za nią kompletnie pustą stronę. Usuwamy je
    i zamiast tego ustawiamy „podział strony przed” na pierwszym akapicie kolejnej
    strony – taki podział nigdy nie tworzy pustej strony, niezależnie od długości tabeli."""
    following = []
    el = table.getnext()
    while el is not None and _is_empty_paragraph(el):
        following.append(el)
        el = el.getnext()

    brk = None          # indeks akapitu z podziałem strony
    brk_before = False  # True: „podział przed akapitem”, False: znak podziału w akapicie
    for i, p in enumerate(following):
        if _has_page_break_before(p):
            brk, brk_before = i, True
        if _has_page_break(p):
            brk, brk_before = i, False
    if brk is None:
        return False

    remove = following[:brk] if brk_before else following[:brk + 1]
    target = following[brk] if brk_before else (following[brk + 1] if brk + 1 < len(following) else el)
    if target is None or target.tag != W + "p":
        return False
    for p in remove:
        p.getparent().remove(p)
    _set_page_break_before(target)
    # Word wymaga akapitu bezpośrednio po tabeli – jeśli następny jest akapit
    # z „podziałem przed”, to jest poprawne; nic więcej nie trzeba.
    return True


# --------------------------------------------------------------------------- #
#  Przetwarzanie
# --------------------------------------------------------------------------- #
def process(docx_path, wb, wb_values, out_path, sheet_name=None, empty_before=DEFAULT_EMPTY_PARAGRAPHS_BEFORE):
    name = sheet_for_document(wb, docx_path, sheet_name)
    document = docx.Document(docx_path)
    anchor, existing = find_insertion_point(document, empty_before)

    widths, rows, _ = read_sheet(wb[name], wb_values[name])
    grid = compute_grid(widths, rows, section_text_width(document, anchor))
    table = parse_xml(build_table_xml(grid, rows))

    if existing is not None:
        existing.addprevious(table)
        existing.getparent().remove(existing)
        action = "podmieniono tabelę"
    else:
        anchor.addnext(table)
        action = "wstawiono tabelę"
    prevent_blank_page(table)
    document.save(out_path)
    return name, action, len(rows)


def missing_formula_values(ws, ws_values):
    """Liczba formuł, dla których plik nie zawiera policzonego wyniku."""
    n = 0
    min_col, min_row, max_col, max_row = sheet_area(ws)
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("=") and \
                    ws_values[c.coordinate].value is None:
                n += 1
    return n


def load_workbooks(excel):
    global THEME_COLORS
    wb = openpyxl.load_workbook(excel)
    wb_values = openpyxl.load_workbook(excel, data_only=True)
    THEME_COLORS = load_theme_colors(wb)
    return wb, wb_values


def list_docx(folder):
    return sorted(f for f in glob.glob(os.path.join(folder, "*.docx"))
                  if not os.path.basename(f).startswith("~$"))


def run_batch(excel, files, out_dir=None, overwrite=False, sheet_name=None,
              empty_before=DEFAULT_EMPTY_PARAGRAPHS_BEFORE, progress=None):
    """Przetwarza listę plików i sprawdza czerwone komórki. Zwraca dane do raportu."""
    wb, wb_values = load_workbooks(excel)
    if out_dir and not overwrite:
        os.makedirs(out_dir, exist_ok=True)

    results = []          # (plik, status, opis)  status: ok / pominiety / blad
    used_sheets = {}      # arkusz -> plik
    for i, f in enumerate(files):
        base = os.path.basename(f)
        if progress:
            progress(i, len(files), base)
        out = f if overwrite else os.path.join(out_dir, base)
        try:
            sheet = sheet_for_document(wb, f, sheet_name)
        except ValueError as e:
            results.append((base, "pominiety", str(e)))
            continue
        try:
            sheet, action, n = process(f, wb, wb_values, out, sheet, empty_before)
            used_sheets[sheet] = base
            results.append((base, "ok", f"{action} z arkusza „{sheet}” ({n} wierszy)"))
        except PermissionError:
            results.append((base, "blad", "nie można zapisać pliku – zamknij go w Wordzie i spróbuj ponownie"))
        except Exception as e:  # noqa: BLE001
            results.append((base, "blad", str(e) or e.__class__.__name__))
    if progress:
        progress(len(files), len(files), "")

    red, no_values = {}, {}
    for ws in wb.worksheets:
        cells = find_red_cells(ws, wb_values[ws.title])
        if cells:
            red[ws.title] = cells
        if ws.title in used_sheets:
            n = missing_formula_values(ws, wb_values[ws.title])
            if n:
                no_values[ws.title] = n
    return {
        "excel": excel, "out_dir": out_dir, "overwrite": overwrite,
        "results": results, "red": red, "used_sheets": used_sheets,
        "no_values": no_values,
        "unused_sheets": [s for s in wb.sheetnames
                          if s not in used_sheets and re.fullmatch(r"\d+", s.strip())],
    }


def _plural_cells(n):
    if n == 1:
        return "1 czerwona komórka"
    if 2 <= n % 10 <= 4 and not 12 <= n % 100 <= 14:
        return f"{n} czerwone komórki"
    return f"{n} czerwonych komórek"


def report_lines(rep):
    """Raport jako lista (tekst, rodzaj); rodzaj: title / ok / warn / err / info."""
    L = []
    ok = sum(1 for r in rep["results"] if r[1] == "ok")
    err = sum(1 for r in rep["results"] if r[1] == "blad")
    skip = sum(1 for r in rep["results"] if r[1] == "pominiety")
    red_docs = {rep["used_sheets"].get(s) for s in rep["red"]}

    L.append(("CZERWONE KOMÓRKI W TABELACH Z OBLICZENIAMI", "title"))
    if rep["red"]:
        for sheet, cells in rep["red"].items():
            doc = rep["used_sheets"].get(sheet)
            where = f"opis {doc}" if doc else "brak opisu w wybranym folderze"
            L.append((f"  ⚠ Arkusz „{sheet}” ({where}) – {_plural_cells(len(cells))}:", "warn"))
            for addr, why, text in cells[:30]:
                text = (text[:90] + "…") if len(text) > 90 else text
                L.append((f"      {addr} – {why}" + (f": „{text}”" if text else ""), "warn"))
            if len(cells) > 30:
                L.append((f"      … i {len(cells) - 30} kolejnych", "warn"))
    else:
        L.append(("  ✔ Brak czerwonych komórek.", "ok"))
    L.append(("", "info"))

    L.append((f"OPISY: wypełniono {ok}, pominięto {skip}, błędów {err}", "title"))
    for base, status, msg in rep["results"]:
        if status == "ok":
            warn = base in red_docs
            L.append((f"  {'⚠' if warn else '✔'} {base} – {msg}"
                      + (" – UWAGA: tabela zawiera czerwone komórki" if warn else ""),
                      "warn" if warn else "ok"))
        elif status == "pominiety":
            L.append((f"  – {base} – pominięto: {msg}", "info"))
        else:
            L.append((f"  ✘ {base} – BŁĄD: {msg}", "err"))
    if rep["no_values"]:
        L.append(("", "info"))
        for sheet, n in rep["no_values"].items():
            L.append((f"  ⚠ Arkusz „{sheet}”: {n} formuł bez policzonej wartości – otwórz plik "
                      f"w Excelu, zapisz go i uruchom program ponownie.", "warn"))
    if rep["unused_sheets"]:
        L.append(("", "info"))
        L.append(("Arkusze bez opisu w wybranym folderze: " + ", ".join(rep["unused_sheets"]), "info"))
    L.append(("", "info"))
    L.append((f"Plik z obliczeniami: {rep['excel']}", "info"))
    L.append(("Zapisano do: " + ("plików oryginalnych" if rep["overwrite"] else str(rep["out_dir"])), "info"))
    return L


def report_text(rep):
    return "\n".join(t for t, _ in report_lines(rep))


# --------------------------------------------------------------------------- #
#  Okienka (tkinter)
# --------------------------------------------------------------------------- #
def run_gui():
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk

    try:  # ostre okienka na monitorach z powiększeniem (Windows)
        import ctypes
        ctypes.windll.shcore.SetProcessDpiAwareness(1)
    except Exception:
        pass

    root = tk.Tk()
    root.title(APP_NAME)
    root.withdraw()
    root.attributes("-topmost", True)

    try:
        excel = filedialog.askopenfilename(
            parent=root, title="Krok 1/3: wskaż plik Excel z obliczeniami",
            filetypes=[("Pliki Excel", "*.xlsx *.xlsm"), ("Wszystkie pliki", "*.*")])
        if not excel:
            return 0

        docs_dir = os.path.dirname(excel)
        while True:
            docs_dir = filedialog.askdirectory(
                parent=root, title="Krok 2/3: wskaż folder z opisami (pliki .docx)",
                initialdir=docs_dir, mustexist=True)
            if not docs_dir:
                return 0
            files = list_docx(docs_dir)
            if files:
                break
            messagebox.showwarning(APP_NAME, f"W folderze\n{docs_dir}\nnie ma plików .docx.\n\n"
                                             "Wskaż inny folder.", parent=root)

        out_dir = filedialog.askdirectory(
            parent=root, title="Krok 3/3: wskaż folder, w którym zapisać gotowe opisy",
            initialdir=docs_dir, mustexist=False)
        if not out_dir:
            return 0
        overwrite = os.path.normcase(os.path.abspath(out_dir)) == os.path.normcase(os.path.abspath(docs_dir))
        if overwrite and not messagebox.askyesno(
                APP_NAME, "Wybrano ten sam folder, w którym są opisy.\n"
                          "Oryginalne pliki zostaną nadpisane.\n\nKontynuować?", parent=root):
            return 0

        win = tk.Toplevel(root)
        win.title(APP_NAME)
        win.resizable(False, False)
        win.attributes("-topmost", True)
        lbl = ttk.Label(win, text="Wczytywanie pliku Excel…", width=60)
        lbl.pack(padx=20, pady=(20, 8))
        bar = ttk.Progressbar(win, length=420, mode="determinate", maximum=max(1, len(files)))
        bar.pack(padx=20, pady=(0, 20))
        win.update()

        def progress(i, n, name):
            bar["value"] = i
            lbl["text"] = f"Przetwarzanie {name} ({i + 1}/{n})…" if name else "Sprawdzanie czerwonych komórek…"
            win.update()

        try:
            rep = run_batch(excel, files, out_dir, overwrite=overwrite, progress=progress)
        finally:
            win.destroy()
        show_report(root, rep)
        return 1 if any(r[1] == "blad" for r in rep["results"]) else 0
    except Exception as e:  # noqa: BLE001
        messagebox.showerror(APP_NAME, f"Wystąpił błąd:\n\n{e}", parent=root)
        return 1
    finally:
        try:
            root.destroy()
        except Exception:
            pass


def show_report(root, rep):
    import tkinter as tk
    from tkinter import ttk

    win = tk.Toplevel(root)
    win.title(f"{APP_NAME} – raport")
    win.geometry("920x620")
    win.attributes("-topmost", True)
    win.after(600, lambda: win.attributes("-topmost", False))

    has_red = bool(rep["red"])
    errors = any(r[1] == "blad" for r in rep["results"])
    head_text = ("UWAGA! W tabelach z obliczeniami są czerwone komórki – sprawdź je."
                 if has_red else "Gotowe. Nie znaleziono czerwonych komórek.")
    if errors:
        head_text += "  Niektórych plików nie udało się zapisać."
    tk.Label(win, text=head_text, font=("Segoe UI", 12, "bold"), fg="white",
             bg="#C00000" if has_red or errors else "#2E7D32", pady=10).pack(fill="x")

    frame = ttk.Frame(win)
    frame.pack(fill="both", expand=True, padx=10, pady=10)
    txt = tk.Text(frame, wrap="word", font=("Segoe UI", 10), relief="flat", padx=8, pady=8)
    sb = ttk.Scrollbar(frame, command=txt.yview)
    txt.configure(yscrollcommand=sb.set)
    sb.pack(side="right", fill="y")
    txt.pack(side="left", fill="both", expand=True)
    txt.tag_configure("title", font=("Segoe UI", 11, "bold"), spacing1=4, spacing3=4)
    txt.tag_configure("ok", foreground="#2E7D32")
    txt.tag_configure("warn", foreground="#C00000")
    txt.tag_configure("err", foreground="#C00000", font=("Segoe UI", 10, "bold"))
    txt.tag_configure("info", foreground="#333333")
    for line, kind in report_lines(rep):
        txt.insert("end", line + "\n", kind)
    txt.configure(state="disabled")

    buttons = ttk.Frame(win)
    buttons.pack(fill="x", padx=10, pady=(0, 10))
    target = rep["out_dir"]
    if target and hasattr(os, "startfile"):
        ttk.Button(buttons, text="Otwórz folder z wynikami",
                   command=lambda: os.startfile(target)).pack(side="left")  # type: ignore[attr-defined]

    def save_report():
        from tkinter import filedialog
        path = filedialog.asksaveasfilename(parent=win, title="Zapisz raport", defaultextension=".txt",
                                            initialdir=target or None, initialfile="raport.txt",
                                            filetypes=[("Plik tekstowy", "*.txt")])
        if path:
            with open(path, "w", encoding="utf-8") as fh:
                fh.write(report_text(rep) + "\n")

    ttk.Button(buttons, text="Zapisz raport…", command=save_report).pack(side="left", padx=8)
    ttk.Button(buttons, text="Zamknij", command=win.destroy).pack(side="right")
    win.protocol("WM_DELETE_WINDOW", win.destroy)
    root.wait_window(win)


# --------------------------------------------------------------------------- #
#  Wiersz poleceń
# --------------------------------------------------------------------------- #
def main(argv=None):
    argv = sys.argv[1:] if argv is None else argv
    if not argv:
        return run_gui()

    ap = argparse.ArgumentParser(
        description="Wstawia tabele obliczeń z pliku Excel do dokumentów Word "
                    "(rozdział „Zestawienie zbiorcze słupów”). Uruchomiony bez "
                    "parametrów otwiera okienka wyboru plików.")
    ap.add_argument("docx", nargs="*", help="pliki .docx lub foldery (domyślnie: bieżący folder)")
    ap.add_argument("-x", "--excel", help="plik Excel z obliczeniami (domyślnie: jedyny .xlsx w folderze)")
    ap.add_argument("-o", "--wynik", default="wynik", help='folder wyjściowy (domyślnie "wynik")')
    ap.add_argument("--nadpisz", action="store_true", help="zapisz zmiany w oryginalnych plikach")
    ap.add_argument("--arkusz", help="nazwa arkusza (gdy nazwa pliku nie odpowiada nazwie arkusza)")
    ap.add_argument("--puste-akapity", type=int, default=DEFAULT_EMPTY_PARAGRAPHS_BEFORE,
                    help="ile pustych akapitów pod nagłówkiem zostawić przed tabelą (domyślnie 5)")
    ap.add_argument("--raport", help="zapisz raport do pliku tekstowego")
    ap.add_argument("--okna", action="store_true", help="otwórz okienka wyboru plików")
    ap.add_argument("--test-okien", help=argparse.SUPPRESS)  # test działania tkinter w .exe
    args = ap.parse_args(argv)

    if args.okna:
        return run_gui()
    if args.test_okien:
        import tkinter as tk
        from tkinter import filedialog, messagebox, ttk  # noqa: F401
        root = tk.Tk()
        root.withdraw()
        root.update()
        root.destroy()
        with open(args.test_okien, "w", encoding="utf-8") as fh:
            fh.write("tk %s OK\n" % tk.TkVersion)
        return 0

    inputs = args.docx or [os.getcwd()]
    files = []
    for p in inputs:
        files += list_docx(p) if os.path.isdir(p) else [p]
    if not files:
        print("Nie znaleziono plików .docx.")
        return 1

    excel = args.excel
    if not excel:
        search_dirs = {os.path.dirname(os.path.abspath(files[0])), os.getcwd()}
        found = sorted({os.path.abspath(f) for d in search_dirs for f in glob.glob(os.path.join(d, "*.xlsx"))
                        if not os.path.basename(f).startswith("~$")})
        if len(found) != 1:
            print("Podaj plik Excel opcją -x (znaleziono: %s)." % (", ".join(found) or "brak"))
            return 1
        excel = found[0]

    rep = run_batch(excel, files, args.wynik, overwrite=args.nadpisz,
                    sheet_name=args.arkusz, empty_before=args.puste_akapity)
    text = report_text(rep)
    print(text)
    if args.raport:
        with open(args.raport, "w", encoding="utf-8") as fh:
            fh.write(text + "\n")
    return 1 if any(r[1] == "blad" for r in rep["results"]) else 0


def _fix_streams():
    # w pliku .exe bez konsoli sys.stdout/stderr nie istnieją
    for name in ("stdout", "stderr"):
        stream = getattr(sys, name)
        if stream is None:
            setattr(sys, name, open(os.devnull, "w", encoding="utf-8"))
        else:
            try:
                stream.reconfigure(errors="replace")
            except Exception:
                pass


if __name__ == "__main__":
    _fix_streams()
    sys.exit(main())
